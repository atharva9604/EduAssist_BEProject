import os
from pptx import Presentation
from pathlib import Path

def inspect_last_generated():
    storage_dir = Path("backend/storage/presentations")
    if not storage_dir.exists():
        print("Storage dir not found")
        return

    # Find newest pptx
    files = list(storage_dir.glob("*.pptx"))
    if not files:
        print("No generated PPTs found")
        return
        
    newest = max(files, key=os.path.getctime)
    print(f"🕵️ Inspecting newest PPT: {newest.name}")
    
    prs = Presentation(newest)
    print(f"Total Slides: {len(prs.slides)}")
    
    for i, slide in enumerate(prs.slides):
        print(f"Slide {i} Shapes: {[s.name for s in slide.shapes]}")
        
    # Check for 'Freeform' shapes (from user's template)
    # We expect them on content slides (1..N-1)
    has_freeform = False
    for i, slide in enumerate(prs.slides):
        if i > 0 and i < len(prs.slides) - 1: # Content slides
            shapes = [s.name for s in slide.shapes]
            if any("Freeform" in s for s in shapes):
                has_freeform = True
                print(f"✅ Slide {i} has Branding: {shapes}")
            else:
                print(f"❌ Slide {i} MISSING Branding: {shapes}")
    
    print(f"\nHas Branding on content slides? {has_freeform}")

if __name__ == "__main__":
    inspect_last_generated()
