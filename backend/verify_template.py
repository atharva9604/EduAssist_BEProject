#!/usr/bin/env python3
import sys
import os
from pathlib import Path
from pptx import Presentation

# Add backend directory to path
backend_dir = Path(__file__).parent
sys.path.insert(0, str(backend_dir))

def create_mock_template():
    """Create a mock template file for testing (3 slides)"""
    # Use absolute path for storage
    backend_root = backend_dir # This is already defined as Path(__file__).parent
    template_dir = backend_root / "storage" / "templates"
    template_dir.mkdir(parents=True, exist_ok=True)
    # CRITICAL: Use a different filename so we don't overwrite user's actual template
    template_path = template_dir / "test_template.pptx"
    
    print(f"Creating mock template at {template_path}")
    prs = Presentation()
    
    # Slide 0: Title Slide
    slide0 = prs.slides.add_slide(prs.slide_layouts[0])
    slide0.shapes.title.text = "TEMPLATE TITLE PLACEHOLDER"
    
    # Slide 1: Content Placeholder (with signature)
    slide1 = prs.slides.add_slide(prs.slide_layouts[1])
    slide1.shapes.title.text = "CONTENT PLACEHOLDER"
    # Add signature to content layout
    left = top = 100
    width = height = 50
    shape = slide1.shapes.add_shape(1, left, top, width, height)
    shape.name = "ContentSignature"
    
    # Slide 2: Thank You Slide
    slide2 = prs.slides.add_slide(prs.slide_layouts[1]) # Use layout 1 or blank
    slide2.shapes.title.text = "THANK YOU SLIDE"
    shape = slide2.shapes.add_shape(1, left, top, width, height)
    shape.name = "ThankYouSignature"
    
    prs.save(str(template_path))
    return template_path

def verify_template_loading():
    """Verify that PPTCreator loads the template and handles 3-slide logic"""
    print("Verifying Template Loading & Logic (3-Slide Mode)...")
    
    try:
        from utils.ppt_creator import PPTCreator
        
        # 1. Create mock template
        template_path = create_mock_template()
        
        # 2. Instantiate Creator with explicit temporary template path handling
        # Since PPTCreator defaults to 'template.pptx', we need to temporarily point it to 'test_template.pptx'
        # OR we just pass the template path if the class supported it. 
        # Looking at PPTCreator, it hardcodes 'template.pptx' in __init__.
        # We should modify it to allow overriding, OR just rename our test file for the test?
        # NO, renaming risks overwriting.
        
        # Let's patch the instance's template_path
        creator = PPTCreator()
        creator.template_path = template_path # Override with "test_template.pptx"
        
        # 3. Create a presentation with multiple content slides
        content_data = {
            "presentation_title": "Template Logic Test",
            "slides": [
                {"slide_number": 1, "slide_type": "title", "title": "My New Title"},
                {"slide_number": 2, "slide_type": "content", "title": "Content 1", "content": ["A"]},
                {"slide_number": 3, "slide_type": "content", "title": "Content 2", "content": ["B"]},
                {"slide_number": 4, "slide_type": "summary", "title": "Summary", "content": ["C"]}
            ]
        }
        
        output_path = creator.create_presentation(content_data, "test_template_logic.pptx")
        print(f"Generated PPT at: {output_path}")
        
        # 4. Inspect the output
        prs = Presentation(output_path)
        
        # CHECK 1: Title Slide (Should be reused, Index 0)
        print(f"Slide 0 Title: {prs.slides[0].shapes.title.text}")
        if prs.slides[0].shapes.title.text == "Template Logic Test": # Title from content_data
             print("✅ Title Slide: Updated correctly.")
        else:
             print(f"❌ Title Slide: Text mismatch. Got '{prs.slides[0].shapes.title.text}'")

        # CHECK 2: Content Slides (Indices 1, 2, 3 should be new content)
        # Verify content slide presence
        print(f"Total Slides: {len(prs.slides)}")
        # Expected: Title(1) + Content1(1) + Content2(1) + Summary(1) + ThankYou(1) = 5 slides total
        # (Original ContentPlaceholder should be removed)
        
        if len(prs.slides) == 5:
             print("✅ Slide Count: Correct (5 slides).")
        else:
             print(f"❌ Slide Count: Incorrect. Expected 5, got {len(prs.slides)}")
             for i, s in enumerate(prs.slides):
                 title = s.shapes.title.text if s.shapes.title else "No Title"
                 print(f"  Slide {i}: {title}")

        # CHECK 3: Thank You Slide (Should be LAST)
        last_slide = prs.slides[-1]
        has_ty_sig = any(shape.name == "ThankYouSignature" for shape in last_slide.shapes)
        
        if has_ty_sig:
            print("✅ Thank You Slide: Is at the END.")
        else:
            print("❌ Thank You Slide: NOT at the end.")
            
        # Clean up
        # os.remove(template_path)
        # os.remove(output_path)

    except Exception as e:
        with open("debug_log.txt", "w") as f:
            f.write(f"Error verification: {e}\n")
            import traceback
            traceback.print_exc(file=f)

if __name__ == "__main__":
    # Redirect stdout to file
    # Use utf-8 encoding explicitly
    with open("debug_log.txt", "w", encoding="utf-8") as f:
        sys.stdout = f
        verify_template_loading()
        sys.stdout = sys.__stdout__
    
    # Print file content to console safely
    try:
        with open("debug_log.txt", "r", encoding="utf-8") as f:
            print(f.read())
    except UnicodeEncodeError:
        print("Could not print log to console due to encoding, please check debug_log.txt")
