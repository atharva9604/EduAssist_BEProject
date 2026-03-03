import copy
from pptx import Presentation

def fix_template():
    input_path = "backend/storage/templates/template.pptx"
    output_path = "backend/storage/templates/template_corrected.pptx"
    
    print(f"🔧 Attempting to fix template: {input_path}")
    prs = Presentation(input_path)
    
    # helper to copy shapes
    def copy_shapes_to_layout(source_slide, target_layout):
        print(f"  Copying shapes from Slide {source_slide.slide_id} to Layout {target_layout.name}...")
        
        # Access XML element trees
        src_spTree = source_slide.shapes._spTree
        tgt_spTree = target_layout.shapes._spTree
        
        count = 0
        for shape in source_slide.shapes:
            # Skip placeholders (we don't want to copy placeholders, we want branding)
            if shape.is_placeholder:
                continue
                
            # Copy the shape element
            new_sp = copy.deepcopy(shape.element)
            
            # Append to target layout
            tgt_spTree.append(new_sp)
            count += 1
            
        print(f"  ✅ Copied {count} shapes.")

    # 1. Fix Title Layout (Layout 0)
    # We assume Slide 0 has the desired Title branding
    if len(prs.slides) > 0:
        slide0 = prs.slides[0]
        layout0 = prs.slide_layouts[0] # Title Slide Layout
        copy_shapes_to_layout(slide0, layout0)
        
    # 2. Fix Content Layout (Layout 1)
    # We assume Slide 1 has the desired Content branding
    if len(prs.slides) > 1:
        slide1 = prs.slides[1]
        layout1 = prs.slide_layouts[1] # Title and Content Layout
        copy_shapes_to_layout(slide1, layout1)
        
    prs.save(output_path)
    print(f"💾 Saved corrected template to: {output_path}")

if __name__ == "__main__":
    try:
        fix_template()
    except Exception as e:
        print(f"❌ Error fixing template: {e}")
        import traceback
        traceback.print_exc()
